"""
AI Document Processor - Portfolio Project
Author: Oscar J. Villa García
Description: Extracts structured data from PDFs, images and text files using Claude API.
             Outputs clean JSON and Excel reports. Zero templates needed.
"""

import os, json, base64, logging, argparse
from pathlib import Path
from datetime import datetime
import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).parent.parent / "output"
CLAUDE_API = "https://api.anthropic.com/v1/messages"
MODEL = "claude-sonnet-4-20250514"
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
TEXT_EXTS  = {".txt", ".md", ".csv", ".json", ".xml", ".html"}
PDF_EXT = ".pdf"

SYSTEM_PROMPT = """You are a document data extraction expert.
Extract ALL structured information from the provided document.
Return ONLY a valid JSON object — no markdown, no explanation.
Use snake_case keys. Group related fields into nested objects.
If a value is not found, use null. Always include a "document_type" field."""

def call_claude(messages, system=""):
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise EnvironmentError("ANTHROPIC_API_KEY not set.")
    payload = {"model": MODEL, "max_tokens": 2048, "messages": messages}
    if system:
        payload["system"] = system
    resp = requests.post(CLAUDE_API, headers={"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"}, json=payload, timeout=60)
    resp.raise_for_status()
    return resp.json()["content"][0]["text"]

def encode_file(path):
    ext = path.suffix.lower()
    media_map = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
    return base64.standard_b64encode(path.read_bytes()).decode("utf-8"), media_map.get(ext, "application/octet-stream")

def build_message(file_path, extraction_hint=""):
    ext = file_path.suffix.lower()
    hint_text = f"\n\nExtraction focus: {extraction_hint}" if extraction_hint else ""
    if ext in IMAGE_EXTS or ext == PDF_EXT:
        data, media_type = encode_file(file_path)
        content_type = "document" if ext == PDF_EXT else "image"
        return [{"role": "user", "content": [{"type": content_type, "source": {"type": "base64", "media_type": media_type, "data": data}}, {"type": "text", "text": f"Extract all structured data from this document.{hint_text}"}]}]
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    return [{"role": "user", "content": f"Extract all structured data from this document:\n\n{text}{hint_text}"}]

def process_file(file_path, extraction_hint=""):
    log.info(f"Processing: {file_path.name}")
    messages = build_message(file_path, extraction_hint)
    raw = call_claude(messages, system=SYSTEM_PROMPT)
    clean = raw.strip().lstrip("```json").lstrip("```").rstrip("```").strip()
    try:
        data = json.loads(clean)
    except json.JSONDecodeError:
        data = {"raw_extraction": raw}
    data["_source_file"] = file_path.name
    data["_processed_at"] = datetime.now().isoformat()
    return data

def flatten(d, prefix=""):
    result = {}
    for k, v in d.items():
        key = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict):
            result.update(flatten(v, key))
        elif isinstance(v, list):
            result[key] = ", ".join(str(i) for i in v)
        else:
            result[key] = v
    return result

def save_outputs(results, label="extraction"):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = OUTPUT_DIR / f"{label}_{ts}.json"
    json_path.write_text(json.dumps(results, indent=2, ensure_ascii=False))
    df = pd.DataFrame([flatten(r) for r in results])
    xlsx_path = OUTPUT_DIR / f"{label}_{ts}.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Extracted Data")
        ws = writer.book.active
        for col in range(1, len(df.columns)+1):
            cell = ws.cell(row=1, column=col)
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(bold=True, color="FFFFFF")
            ws.column_dimensions[get_column_letter(col)].width = 22
        ws.freeze_panes = "A2"
    log.info(f"Saved: {json_path} | {xlsx_path}")
    return {"json": json_path, "xlsx": xlsx_path}

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AI Document Processor — Claude API")
    parser.add_argument("input", help="Path to file or directory")
    parser.add_argument("--hint", default="", help="Extraction focus hint")
    parser.add_argument("--label", default="extraction")
    args = parser.parse_args()
    target = Path(args.input).resolve()
    results = [process_file(f, args.hint) for f in target.iterdir() if f.is_file()] if target.is_dir() else [process_file(target, args.hint)]
    paths = save_outputs(results, args.label)
    print(f"\n✅ Processed {len(results)} document(s)\n   JSON  -> {paths['json']}\n   Excel -> {paths['xlsx']}")
